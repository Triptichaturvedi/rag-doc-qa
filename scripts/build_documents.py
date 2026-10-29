"""Build reports/Project_Report.docx, reports/Project_Synopsis.docx and reports/Ethical_Compliance_Declaration.docx.

All numbers are read from reports/eval_results_tfidf.json and data/cleaning_report.json so documents match the code.
Usage: python -m scripts.build_documents
"""
import json
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
REP, DOCS, FIG = ROOT / "reports", ROOT / "docs", ROOT / "reports" / "figures"
EV = json.loads((REP / "eval_results_tfidf.json").read_text())
CL = json.loads((ROOT / "data" / "cleaning_report.json").read_text())
G, S, AB = EV["gold"], EV["synthetic_cloze"], EV["abstention"]
DF = pd.read_csv(ROOT / "data" / "dataset.csv")
pct = lambda x, d=1: f"{x * 100:.{d}f}%"  # noqa: E731

FILL = "[FILL IN]"
AI_DISCLOSURE = (
    "AI assistance disclosure: the source code, corpus text, evaluation scripts, notebook and the first drafts of "
    "these documents were produced with the assistance of Claude (Anthropic's AI assistant) in response to the learner's "
    "instructions. Scheme facts that change over time were checked against current web sources on 28 October 2026 "
    "(for example the replacement of MGNREGA by VB-G RAM G on 1 July 2026 and NIST SP 800-63B Revision 4). The learner is "
    "responsible for reviewing, running and understanding all submitted material before submission."
)


# ----------------------------------------------------------------------------- docx helpers
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_cell(cell, text, bold=False, size=9, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align:
        p.alignment = align


def add_page_number(section):
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("RAG-Based Document Q&A System  |  Page ")
    run.font.size = Pt(8)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run(); r.font.size = Pt(8)
        if tag:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tag); r._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2.2)
        st = self.d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
        for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
            h = self.d.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        add_page_number(s)
        self.fig_n = self.tab_n = 0

    def h(self, text, level=1):
        return self.d.add_heading(text, level)

    def p(self, text="", bold=False, italic=False, size=None, align=None, after=None):
        para = self.d.add_paragraph()
        self._runs(para, text, bold, italic, size)
        if align:
            para.alignment = align
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    @staticmethod
    def _runs(para, text, bold=False, italic=False, size=None):
        # **bold** inline markup
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = para.add_run(part); r.bold = bold or (i % 2 == 1); r.italic = italic
            if size:
                r.font.size = Pt(size)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style); self._runs(para, it); para.paragraph_format.space_after = Pt(2)

    def table(self, header, rows, widths=None, size=9, caption=None):
        if caption:
            self.tab_n += 1
            c = self.p(f"Table {self.tab_n}. {caption}", bold=True, size=9.5, after=3); c.paragraph_format.keep_with_next = True
        t = self.d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(header):
            set_cell(t.rows[0].cells[i], htxt, bold=True, size=size, color="FFFFFF"); shade(t.rows[0].cells[i], "1F4E79")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                set_cell(cells[i], v, size=size)
                if ri % 2 == 1:
                    shade(cells[i], "EAF1F8")
        if widths:
            t.autofit = False
            for i, w in enumerate(widths):
                t.columns[i].width = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        for ri in (0, 1):                                   # keep header + first row together with what follows
            if ri < len(t.rows):
                for c in t.rows[ri].cells:
                    for para in c.paragraphs:
                        para.paragraph_format.keep_with_next = True
        for row in t.rows:                                  # never split a row across pages
            trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def fig(self, path, caption, width=15.5):
        self.d.add_picture(str(path), width=Cm(width))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.fig_n += 1
        self.p(f"Figure {self.fig_n}. {caption}", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.d.save(str(path)); print("wrote", path)


# ----------------------------------------------------------------------------- REPORT
def build_report():
    D = Doc()
    # ---- cover
    for _ in range(3):
        D.p("")
    D.p("HCLTech Career Shaper Programme", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    D.p("III Year Industry Project  |  Academic Year 2025-26", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    D.p("RAG-Based Document Q&A System", bold=True, size=26, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    D.p("Project No. 72  |  Domain: LLM / Python  |  Advanced", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    D.p("Project Report", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    D.table(["Field", "Details"], [["Learner name", FILL], ["Roll number", FILL], ["College / campus", FILL],
                                   ["Mentor", FILL], ["Submission date", FILL], ["Format", "Individual project (1 learner), 4 weeks"]],
            widths=[5, 10], size=10)
    D.p("")
    D.p("Retrieval-Augmented Generation for question answering over public-domain documents (Indian Constitution, government "
        "schemes, technical manuals) using LangChain, ChromaDB and an OpenAI / open-source LLM, with source display, "
        "abstention and legal/medical safeguards.", italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    D.page_break()

    # ---- abstract + contents
    D.h("Abstract", 1)
    D.p(f"This report describes the design, implementation and evaluation of a Retrieval-Augmented Generation (RAG) system that answers "
        f"questions over a corpus of 50 public documents: 20 on the Constitution of India, 15 on government schemes and Acts, and 15 "
        f"technical manuals from US federal agencies. The system splits documents into sentence-aware chunks, retrieves the most similar "
        f"chunks, and returns an answer together with the exact source chunks, a retrieval-confidence score and, when appropriate, a "
        f"legal or medical disclaimer. If the best chunk is not similar enough, it abstains instead of guessing. The code supports a "
        f"ChromaDB vector store through LangChain with local or OpenAI embeddings, an OpenAI or Ollama chat model, and a fully offline "
        f"TF-IDF + extractive baseline that needs no API key.")
    D.p(f"On 150 hand-written gold questions the offline baseline retrieves the correct document first {pct(G['doc_hit@1'])} of the time "
        f"(hit@3 {pct(G['doc_hit@3'])}, MRR {G['doc_mrr@5']:.3f}), returns an answer containing the gold span in {pct(G['answer_correct_rate'])} of cases, "
        f"and refuses {pct(AB['out_of_scope_correctly_abstained'])} of 30 out-of-scope questions while wrongly refusing {pct(AB['in_scope_wrongly_abstained'])} of "
        f"in-scope ones. A 5,000-row monitoring dataset was cleaned, profiled in an EDA notebook with six visualisations, and documented in a "
        f"data dictionary. **Honest scope note:** the dense-embedding and LLM configurations are implemented but were not executed in the "
        f"development environment (no network access), so no measured results are claimed for them; Section 7.6 lists the commands to produce them.")
    D.h("Contents", 2)
    D.bullets(["1 Introduction and problem statement", "2 Objectives and scope", "3 Background and related work", "4 Data: corpus, Q&A pairs and monitoring dataset",
               "5 Exploratory data analysis", "6 System design and implementation", "7 Evaluation and results", "8 Testing", "9 Ethical considerations and compliance",
               "10 Deployment", "11 Challenges and resolutions", "12 Limitations and future work", "13 Conclusion", "References", "Appendices A-D"])
    D.page_break()

    # ---- 1 Intro
    D.h("1. Introduction and problem statement", 1)
    D.p("**Who** - students, analysts and citizens who must find specific facts in long public documents. **What** - looking up a figure, an article number or "
        "a procedure in a 300-page PDF is slow, and generic chatbots answer fluently but can invent details (hallucination). **When/where** - whenever "
        "someone needs a quick, checkable answer, for example the benefit amount of a welfare scheme or the minimum password length in a security standard. "
        "**Why it matters** - a wrong figure or a wrong article can mislead decisions, and in legal or health-related contexts it can cause harm. "
        "**How** - a RAG system grounds every answer in retrieved passages, shows those passages to the user, and abstains when the evidence is weak.")
    D.p("The problem is concrete. In the corpus built for this project a fact such as \"MGNREGA guaranteed 100 days of work\" became out of date on 1 July 2026, "
        "when the Viksit Bharat G RAM G Act replaced it with a 125-day guarantee. A system that answers from memory would silently repeat the old figure; "
        "a RAG system that cites its source, carries an \"as of\" date and lets the owner update the documents can be corrected in minutes.")
    D.h("2. Objectives and scope", 1)
    D.bullets([
        "Build a working prototype (Streamlit) that answers questions over PDF/Word/TXT/Markdown documents using LangChain and ChromaDB, with an OpenAI or open-source (Ollama) LLM option.",
        "Create a corpus of 50 public-domain/open-licensed documents and 1,000 Q&A pairs, plus a cleaned 5,000-row monitoring dataset with a complete data dictionary.",
        "Measure retrieval quality, answer correctness, abstention behaviour and latency on held-out questions, and report failures honestly.",
        "Implement the ethical requirements of the brief: always display the source chunk beside the answer; do not give legal or medical advice without expert review.",
        "Package the project so that it can be pushed to GitHub and deployed (Streamlit Community Cloud) with a documented fallback that needs no API key."])
    D.p("**Out of scope:** fine-tuning models, multi-lingual retrieval, OCR of scanned PDFs, user authentication, and certified legal or medical advice.")

    # ---- 3 Background
    D.h("3. Background and related work", 1)
    D.p("Retrieval-Augmented Generation combines a non-parametric memory (a document index) with a generative model, so the model can condition its answer on retrieved "
        "evidence rather than only on its training data [1]. Retrieval can be sparse, such as TF-IDF weighting [2], or dense, where a neural encoder maps queries and passages "
        "to vectors and similarity search finds the nearest ones [3]. Sentence encoders such as the MiniLM models distributed with Sentence-Transformers make dense retrieval "
        "cheap enough to run on a CPU [4]. Language models can still produce fluent but unsupported statements, which is the hallucination problem surveyed in [5], and "
        "they may under-use information placed in the middle of long contexts [6]. The NIST AI Risk Management Framework recommends documenting limitations, measuring "
        "risks and keeping humans in the loop [7], and India's Digital Personal Data Protection Act, 2023 sets duties for processing personal data [8].")
    D.p("These sources shaped four design decisions: (i) always return source chunks, (ii) keep the number of retrieved passages small (top-4) to limit context dilution, "
        "(iii) calibrate an abstention threshold instead of trusting the model, and (iv) retain a no-LLM extractive mode that cannot hallucinate because it only copies text.")

    # ---- 4 Data
    D.h("4. Data: corpus, Q&A pairs and monitoring dataset", 1)
    D.h("4.1 Document corpus (50 documents)", 2)
    D.p("The corpus consists of 50 documents of about 175 words each (8,723 words in total). Because the sandbox used to build the project had no internet access, "
        "the documents are condensed, original-wording paraphrases of public material, each carrying front-matter metadata (doc_id, title, category, source, source_url, as_of). "
        "They are not verbatim copies of the originals. Time-sensitive scheme facts were verified against current web sources on 28 October 2026 before writing.")
    D.table(["Category", "Docs", "Examples", "Source / licence note"],
            [["Constitution of India", "20", "Preamble; Fundamental Rights; Directive Principles; President; Parliament; Supreme Court; Emergency; Article 368 and basic structure; Panchayati Raj",
              "Government of India (Legislative Department); reuse permitted with attribution"],
             ["Government schemes and Acts", "15", "PM-KISAN; Ayushman Bharat PM-JAY; MGNREGA/VB-G RAM G; PMJDY; PMAY; Jal Jeevan Mission; RTI Act; NFSA; RTE Act",
              "Government of India ministries and PIB (GODL-India)"],
             ["Technical manuals", "15", "NIST CSF 2.0; NIST SP 800-63B; SP 800-61; SP 800-88; AI RMF; SP 800-53; NASA SE Handbook; FEMA; Army FM 3-25.26; OSHA LOTO and HazCom; USDA temperatures",
              "US federal works (public domain)"]],
            widths=[3.2, 1.2, 7.2, 4.6], caption="Corpus composition")
    D.h("4.2 Q&A pairs (1,000)", 2)
    D.p("The Q&A set has 1,000 pairs: **150 hand-written gold pairs** (three per document, covering the five query types factual, procedural, comparative, analytical and "
        "summarization) and **850 synthetic cloze pairs** generated automatically by masking a number, article reference, named entity or rare technical term in a source sentence "
        "(20 pairs per document in total). Every answer is an exact span of its document, which a unit test verifies. A further 30 out-of-scope questions (for example "
        "\"What is the capital of France?\") test abstention. The synthetic pairs quote nearly verbatim text and the document title, so they are a sanity check rather than a benchmark; "
        "headline results use only the gold pairs.")
    D.h("4.3 Monitoring dataset (dataset.csv)", 2)
    D.p(f"The brief supplies dataset.csv with seven columns (doc_id, doc_type, chunk_id, word_count, retrieval_score, query_type, answer_quality_score) and 500 rows. "
        f"It was extended to {CL['rows_out']:,} rows with the same columns and value ranges (retrieval 0.5-1.0, quality 0.6-1.0, 50-499 words). The first 500 rows are untouched. "
        f"About 3% of the new rows are injected faults of three kinds: retrieval failure (retrieval below 0.5), hallucination suspect (retrieval at least 0.85 but quality below 0.55) and "
        f"truncated chunk (fewer than 50 words). The data are synthetic, so the scores do not measure a real system. The brief's timestamp rule does not apply because "
        f"dataset.csv has no timestamp column.")
    D.p("**Cleaning and pre-processing** (scripts/prepare_dataset.py): whitespace and case normalisation; type validation; null and duplicate checks "
        f"(none found); range checks; harmonising doc_type so that each document has exactly one format ({CL['doc_type_values_changed']:,} of 5,000 values changed, because the raw data "
        f"listed up to five different formats for the same document); rule-based anomaly flags ({CL['anomaly_rate'] * 100:.2f}% of rows, {sum(v for k, v in CL['anomalies'].items() if k != 'none')} rows, "
        "which exactly match the injected labels); and derived columns (category, anomaly_flag, anomaly_type, chunk_length_band, retrieval_band, quality_band). Anomalies are retained and flagged, "
        "not deleted. data_dictionary.csv describes all 14 columns.")

    # ---- 5 EDA
    D.h("5. Exploratory data analysis", 1)
    D.p("The full analysis is in notebooks/EDA.ipynb (13 executed cells, six visualisations, statistical tests). The main findings are summarised here.")
    D.fig(FIG / "eda_distributions.png", "Distributions of retrieval score and answer-quality score; anomalies (red) lie outside the normal ranges.", 15.5)
    D.p("Normal rows are close to uniform inside the documented ranges (retrieval mean 0.750, quality mean 0.800). Anomalous rows have a much lower mean quality "
        f"(0.367) and, for retrieval failures, retrieval scores down to 0.05. Because the anomalies break the normal ranges, simple threshold rules detect them exactly.")
    D.fig(FIG / "eda_scatter_corr.png", "Left: retrieval score versus answer quality with anomaly types highlighted. Right: Spearman correlations.", 15.5)
    D.p("Among normal rows retrieval score and answer quality are uncorrelated (Spearman rho = -0.016, p = 0.26). This is a property of the synthetic generator and "
        "**must not be read as a property of real RAG systems**, where better retrieval normally improves answers. Hallucination suspects form a distinct cluster at "
        "high retrieval and low quality, the pattern a monitoring dashboard would alert on.")
    D.fig(FIG / "eda_query_type.png", "Query-type frequencies and mean answer quality by query type (normal rows, 95% CI).", 15.0)
    D.p("Quality does not differ significantly between query types (Kruskal-Wallis H = 4.44, p = 0.350); chunk length does not differ by document type (p = 0.886); and the anomaly "
        "rate does not depend on the document category (chi-square p = 0.814). These null results are expected for generated data and document that no group is accidentally privileged.")
    D.fig(FIG / "eda_anomaly.png", "Anomaly rate (%) by category, document type and query type; dashed line is the overall rate.", 15.5)
    D.fig(FIG / "eda_real_eval.png", "Measurements from the real pipeline (not synthetic): top-1 score by retrieval correctness, and answer correctness by query type and category.", 15.5)
    D.p("Section 8 of the notebook switches from synthetic to real data. The top-1 retrieval score separates in-scope questions (median "
        f"{AB['median_top_score_in_scope']:.2f}) from out-of-scope questions (median {AB['median_top_score_out_of_scope']:.2f}), which is what makes score-based abstention possible. "
        "Only five gold questions retrieved a wrong document first, too few to say whether the score also predicts retrieval errors.")

    # ---- 6 Design
    D.h("6. System design and implementation", 1)
    D.h("6.1 Architecture", 2)
    D.fig(DOCS / "architecture_dfd0.png", "Data Flow Diagram, Level 0 (context diagram).", 13.5)
    D.fig(DOCS / "architecture_blocks.png", "System block diagram: ingestion, question answering and evaluation paths.", 16.0)
    D.p("Ingestion loads documents (Markdown with metadata, plus PDF through pypdf and Word through python-docx for uploads), cleans whitespace, and splits text into "
        "sentence-aware chunks of about 90 words with one sentence of overlap, producing 155 chunks for the 50-document corpus. The retriever returns the top-k chunks "
        "with a similarity score in [0, 1]. The pipeline classifies the query type, checks for legal/medical advice triggers, abstains if the best score is below the threshold, and otherwise "
        "generates the answer. The user interface always shows the source chunks next to the answer.")
    D.h("6.2 Components", 2)
    D.table(["Module", "Responsibility"],
            [["ragqa/ingest.py", "Loaders (md, pdf, docx, txt), sentence splitter, chunker with overlap; Chunk and Doc data classes"],
             ["ragqa/retrieval.py", "TfidfRetriever (word + character n-gram blend, offline); ChromaRetriever (LangChain Chroma, cosine); local ONNX MiniLM or OpenAI embeddings; SQLite shim for Streamlit Cloud"],
             ["ragqa/generation.py", "ExtractiveAnswerer (IDF-weighted sentence selection, cannot hallucinate); LLMAnswerer (LangChain prompt | chat model | parser for OpenAI or Ollama)"],
             ["ragqa/guardrails.py", "Abstention rule; legal and medical advice detectors with notices"],
             ["ragqa/pipeline.py", "RAGPipeline.ask(): retrieve, guard, generate, fall back to extractive if the LLM is unavailable, return Answer with sources"],
             ["ragqa/evaluation.py", "Hit@k, MRR, support and answer recall, false-abstain rate, threshold calibration"],
             ["ragqa/telemetry.py", "Interaction log in the dataset.csv schema plus question, answer and feedback"],
             ["app.py", "Streamlit UI: settings, uploads, chat, source panel, feedback, corpus browser, evaluation tab"],
             ["scripts/", "build_corpus, make_qa, prepare_dataset, evaluate, build_notebook, make_diagrams, build_documents"]],
            widths=[4.0, 12.2], caption="Code modules")
    D.h("6.3 Key design decisions", 2)
    D.bullets([
        "**Three retrieval backends behind one interface** (tfidf, chroma-onnx, chroma-openai) so the project runs anywhere and the dense options can be compared on identical questions.",
        "**Sentence-aware chunking** keeps facts intact; a fixed-size character splitter would cut sentences such as \"Rs 6,000 per year in three equal instalments\" in half.",
        "**Prompt contract for the LLM:** answer only from numbered context, cite passages like [1], and reply exactly \"I don't know based on the provided documents.\" when the context is insufficient; temperature 0.",
        "**Graceful degradation:** if an API key is missing or a package is not installed, the app explains why and continues with TF-IDF + extractive instead of crashing.",
        "**Extractive default:** the safe default only copies text from retrieved chunks, so it cannot fabricate a figure; LLM generation is opt-in.",
        "**Word + character TF-IDF blend:** the character n-grams tolerate hyphenation and plurals (\"PM-KISAN\"); this was added after a failure analysis (Section 11)."])
    D.h("6.4 User interface", 2)
    D.p("The Ask tab shows the answer, the top retrieval score, query type, latency and generator. A bordered panel lists every source chunk with title, doc_id, chunk_id, "
        "category, format and score, expanded by default so the evidence is never hidden. Legal or medical notices appear as an info box. Thumbs-up and thumbs-down buttons log feedback. "
        "The Corpus tab lists the 50 documents; the Evaluation tab shows the measured metrics and figures. The sidebar lets the user choose the backend and generator, chunk size, top-k and "
        "abstention threshold, and upload their own PDF, Word, TXT or Markdown files.")

    # ---- 7 Evaluation
    D.h("7. Evaluation and results", 1)
    D.h("7.1 Method", 2)
    D.p("All metrics come from scripts/evaluate.py, which runs the pipeline on the 1,000 Q&A pairs and 30 out-of-scope questions and writes eval_results_tfidf.json. "
        "**Doc hit@k**: the gold document appears among the top-k chunks. **MRR@5**: mean reciprocal rank of the gold document. **Answer correct**: at least 80% of the content tokens of the gold span "
        "appear in the returned answer (abstentions count as incorrect). **Abstention**: share of out-of-scope questions whose top score is below the threshold, and share of in-scope questions wrongly refused. "
        f"Configuration: offline TF-IDF retriever, extractive answerer, 155 chunks of about 90 words, top-4, abstention threshold {AB['threshold']:.2f}.")
    D.h("7.2 Retrieval and answer quality", 2)
    rows = []
    for name, m in (("Gold (hand-written)", G), ("Synthetic cloze", S)):
        rows.append([name, m["n"], pct(m["doc_hit@1"]), pct(m["doc_hit@3"]), f"{m['doc_mrr@5']:.3f}", pct(m["support_hit@3"]), pct(m["answer_correct_rate"]), pct(m["false_abstain_rate"])])
    D.table(["Question set", "n", "Doc hit@1", "Doc hit@3", "MRR@5", "Support hit@3", "Answer correct", "False abstain"], rows,
            widths=[3.6, 1.0, 1.8, 1.8, 1.6, 2.2, 2.2, 2.0], caption="Results of the offline TF-IDF + extractive baseline")
    D.p(f"On the gold set the right document is ranked first for {pct(G['doc_hit@1'])} of questions and within the top three for {pct(G['doc_hit@3'])}. The synthetic set scores "
        f"{pct(S['doc_hit@1'])} hit@1 and {pct(S['answer_correct_rate'])} answer-correct because its questions copy the wording of the source sentence and include the document title; "
        "it confirms that the plumbing works but says little about real difficulty. The gold figures are the ones to quote. The corpus has only 50 documents, so retrieval is an easier problem "
        "than on a large enterprise collection and the numbers will not transfer unchanged.")
    D.h("7.3 Abstention (out-of-scope questions)", 2)
    D.table(["Metric", "Value"],
            [["Median top-1 score, in-scope gold questions", f"{AB['median_top_score_in_scope']:.3f}"],
             ["Median top-1 score, out-of-scope questions", f"{AB['median_top_score_out_of_scope']:.3f}"],
             [f"Out-of-scope questions correctly refused (threshold {AB['threshold']:.2f}, n = {AB['n_out_of_scope']})", pct(AB["out_of_scope_correctly_abstained"])],
             [f"In-scope questions wrongly refused (n = {AB['n_in_scope']})", pct(AB["in_scope_wrongly_abstained"])],
             ["Threshold maximising balanced accuracy", f"{EV['config']['calibrated_threshold']:.2f}"]],
            widths=[11.5, 4.5], caption="Abstention behaviour")
    D.fig(FIG / "abstention_curve_tfidf.png", "Trade-off between refusing out-of-scope questions and wrongly refusing in-scope ones as the threshold changes.", 11.5)
    D.p("The default threshold of 0.15 deliberately favours answering over refusing (balanced accuracy is highest at 0.19). Because the extractive mode only copies text, a "
        "wrong-but-answered question costs less than it would with free generation. With an LLM the threshold should be raised. Near-domain questions that are not in the corpus "
        "(for example \"What is Article 370?\") are the hardest to refuse because they share vocabulary with the documents. Of the seven out-of-scope questions that were answered, four were near-domain "
        "(Article 370, the Atal Pension premium, the Sukanya Samriddhi rate, 2019 Lok Sabha seats) and three shared an incidental word with a document (\"Prime Minister\", \"population\").")
    D.h("7.4 Breakdown by corpus category and query type", 2)
    rows = [[r["category"], r["n"], pct(r["doc_hit@1"]), pct(r["doc_hit@3"]), pct(r["answer_correct_rate"])] for r in EV["gold_by_category"]]
    rows += [[f"{r['query_type']} (query type)", r["n"], pct(r["doc_hit@1"]), pct(r["doc_hit@3"]), pct(r["answer_correct_rate"])] for r in EV["gold_by_query_type"]]
    D.table(["Group (gold set)", "n", "Doc hit@1", "Doc hit@3", "Answer correct"], rows, widths=[6.0, 1.5, 2.8, 2.8, 3.2], caption="Gold-set results by group")
    D.p("Technical manuals are retrieved best, probably because their vocabulary is distinctive. Comparative questions have the lowest answer correctness (62.5%, n = 8) because "
        "a sentence-level extractor cannot combine two facts; generative or multi-sentence answers are needed. Groups with n below 20 should be read with caution.")
    D.h("7.5 Chunk-size sensitivity and latency", 2)
    rows = [[r["chunk_words"], r["n_chunks"], pct(r["doc_hit@1"]), pct(r["doc_hit@3"]), pct(r["answer_correct_rate"])] for r in EV["chunk_sensitivity_gold"]]
    D.table(["Chunk size (words)", "Chunks", "Doc hit@1", "Doc hit@3", "Answer correct"], rows, widths=[3.6, 2.4, 3.0, 3.0, 3.4],
            caption="Sensitivity of the gold-set results to chunk size (reported, not used to choose the default)")
    D.p(f"The default of 90 words was fixed in advance; the sweep shows the results are stable, with larger chunks slightly improving document ranking and smaller chunks slightly improving "
        f"sentence-level answers. Mean end-to-end latency of retrieval plus extractive answering is about {G['latency_ms_mean']:.1f} ms per question on a CPU (95th percentile {G['latency_ms_p95']:.1f} ms); "
        "an LLM call would add network and generation time that was not measured.")
    D.h("7.6 Dense retrieval and LLM generation (not yet measured)", 2)
    D.p("The ChromaDB, MiniLM, OpenAI and Ollama code paths are implemented but could not be executed or evaluated in the development environment, which had no network access and no "
        "LangChain or ChromaDB installation. Their results must be produced on a networked machine; the evaluation script writes the same files for each configuration so they can be compared directly:")
    D.table(["Configuration", "Command", "Doc hit@1", "Answer correct", "Status"],
            [["TF-IDF + extractive (baseline)", "python -m scripts.evaluate --sweep", pct(G["doc_hit@1"]), pct(G["answer_correct_rate"]), "Measured"],
             ["Chroma + MiniLM + extractive", "python -m scripts.evaluate --backend chroma-onnx", "-", "-", "To be run by learner"],
             ["Chroma + OpenAI embeddings + OpenAI LLM", "python -m scripts.evaluate --backend chroma-openai --llm openai", "-", "-", "To be run by learner"]],
            widths=[4.6, 6.2, 1.8, 2.0, 2.4], size=8.5, caption="Evaluation plan for the remaining configurations")
    D.p("Before relying on the dense configurations, re-calibrate the abstention threshold (the script reports the optimum) because cosine scores from neural embeddings have a different scale than TF-IDF scores; "
        "the app's default of 0.30 for them is only a starting point.")
    D.h("7.7 Failure analysis", 2)
    D.p("Of the 150 gold questions, 30 did not produce an answer containing the gold span. They fall into three groups (five questions retrieved a wrong document first; one of them was refused and counted under false abstention, and one still returned a span-bearing answer):")
    D.table(["Failure type", "Count", "Example", "Cause"],
            [["False abstention", "10", "\"What does Article 24 prohibit?\" (top score 0.148)", "Short questions with few distinctive words score just below the 0.15 threshold"],
             ["Right document, wrong sentence", "17", "\"How many days of work does VB-G RAM G guarantee?\" returned the 15-day work-on-demand sentence instead of the 125-day guarantee",
              "Lexical overlap rewards a different sentence containing \"days\" and \"work\"; no semantic understanding"],
             ["Wrong document ranked first", "3", "\"Which article protects life and personal liberty?\" retrieved the Emergency document", "Article 21 is discussed in several documents; lexical retrieval cannot tell which is the main one"]],
            widths=[3.4, 1.4, 6.4, 5.0], size=8.5, caption="Failure analysis on the gold set (counts add to 30)")
    D.p("Most errors are in answer selection rather than retrieval, which is the expected strength of an LLM generator reading the retrieved passages. This is the main reason to run the LLM configuration, "
        "provided its answers are checked against the displayed source chunks.")

    # ---- 8 Testing
    D.h("8. Testing", 1)
    D.p("Fourteen automated tests (tests/) pass. They check that the corpus has 50 documents in the 20/15/15 split; that the Q&A file has 1,000 unique pairs covering all 50 documents and that every answer "
        "is an exact span of its document; that dataset.csv has the required schema, valid ranges, an anomaly rate between 2% and 5% and one format per document; that the first 500 rows are unchanged; "
        "that the data dictionary has no placeholder text; chunking covers all text and respects the size limit; every answer, even an abstention, ships source chunks; a known question is answered from the right document; "
        "an out-of-scope question abstains; gold hit@3 stays at or above 95%; the legal/medical detectors fire and stay quiet on neutral questions; and the query-type classifier labels five sample questions correctly. "
        "A continuous-integration workflow (.github/workflows/ci.yml) runs the tests and the evaluation on every push.")
    D.p("The Streamlit app was executed end to end against a stub of the Streamlit API to catch logic and name errors (question, answer rendering, feedback and re-run), and compiled cleanly. "
        "It has **not** been run in a real browser session, so the visual layout and the Chroma/LLM paths still need a manual smoke test (Section 10).")

    # ---- 9 Ethics
    D.h("9. Ethical considerations and compliance", 1)
    D.table(["Risk or duty", "How the project addresses it"],
            [["Hallucination (brief: always display the source chunk)", "Every answer is accompanied by its source chunks, expanded by default, with document, chunk id and score. Abstention when evidence is weak. Extractive mode cannot fabricate. LLM prompt forbids unsupported figures."],
             ["Legal or medical advice (brief: not without expert review)", "Regex triggers add a notice for personal-advice questions; the app banner states the limitation; scheme and Constitution texts are informational summaries, and the PM-JAY and food-safety documents carry 'not medical advice' wording."],
             ["Out-of-date information", "Each document has an as_of date. MGNREGA's replacement on 1 July 2026 shows how quickly facts change; users are told to verify against the official portal."],
             ["Licence and provenance", "Only public government material is used (Government of India publications; US federal works in the public domain). Texts are paraphrases with the source and URL recorded in corpus_manifest.csv. The code is MIT-licensed; the licence note excludes the underlying publications."],
             ["Synthetic data", "The monitoring dataset and the 850 cloze questions are synthetic and labelled as such everywhere; no conclusions about real users are drawn from them."],
             ["Personal data (DPDP Act 2023)", "No personal data is in the corpus or datasets. Typed questions are logged only to a local, git-ignored CSV; users are advised not to type personal data. Uploaded files are processed in memory. If an OpenAI option is chosen, text is sent to that provider, which the user must accept."],
             ["Bias and coverage", "The corpus covers 50 documents and mostly English formal text; results do not generalise to other languages or domains."],
             ["Transparency", "AI assistance in producing the code and documents is disclosed in Appendix A and in the Ethical Compliance Declaration."]],
            widths=[4.6, 11.6], size=9, caption="Ethical risks and mitigations")

    # ---- 10 Deployment
    D.h("10. Deployment", 1)
    D.p("The repository is ready for Streamlit Community Cloud. After pushing to GitHub, create a new app on share.streamlit.io, select the repository and branch, set the main file to app.py and choose "
        "Python 3.11 in the advanced settings. Add OPENAI_API_KEY under Secrets only if OpenAI options are wanted. requirements.txt installs the full stack; if the build fails on chromadb, copy "
        "requirements-lite.txt over requirements.txt and the app runs in offline TF-IDF mode with identical features except dense retrieval. The README gives the Git commands and a post-deploy smoke-test checklist: ask an example question, "
        "confirm the source chunks appear, ask an out-of-scope question and confirm abstention, upload a small PDF, and switch backends.")
    D.p("Operational notes: the ChromaDB index and interaction log are written to the container's local disk, which Streamlit Cloud resets on restart, so nothing is persisted; free-tier memory is limited, "
        "which is why the default embedding option avoids PyTorch; and API keys must never be committed (.streamlit/secrets.toml is git-ignored).")

    # ---- 11 Challenges
    D.h("11. Challenges and resolutions", 1)
    D.table(["Challenge", "Resolution"],
            [["Facts in the corpus can go stale: from memory, MGNREGA would still have been described with a 100-day guarantee.", "Checked volatile schemes and the NIST password standard against current sources; rewrote the MGNREGA document to cover its 2026 replacement; added as_of metadata and a user warning."],
             ["The raw dataset gave the same document up to five different formats and had no timestamp or label columns.", "Harmonised doc_type per document, kept the original in doc_type_original, derived anomaly flags by rule and checked them against the label file; documented that the timestamp rule is not applicable."],
             ["Lexical retrieval matched \"money\" in a question about PM-KISAN to a document titled \"Money Bills\".", "Blended word and character n-gram TF-IDF (gold hit@1 rose from 92.7% to 96.7%) and down-weighted title words when choosing answer sentences. This change was made after seeing the gold-set failures, so the gold results are slightly optimistic."],
             ["Synthetic cloze questions gave near-perfect scores that hid real difficulty.", "Reported gold and synthetic results separately and based all headline claims on the 150 hand-written questions."],
             ["No network access or LangChain/ChromaDB in the build environment.", "Used lazy imports and a pure scikit-learn baseline so that the whole pipeline is testable offline; clearly marked the dense and LLM paths as untested."],
             ["ChromaDB needs a newer SQLite than Streamlit Community Cloud provides.", "Added a pysqlite3-binary shim and a lite requirements fallback."]],
            widths=[7.0, 9.2], size=9, caption="Challenges and how they were resolved")

    # ---- 12 Limitations
    D.h("12. Limitations and future work", 1)
    D.bullets([
        "**Untested paths:** dense retrieval, OpenAI/Ollama generation and the live Streamlit UI were not run here (Sections 7.6 and 8).",
        "**Small corpus:** 50 documents, about 175 words each; a real enterprise collection has far more near-duplicates and longer documents, so retrieval will be harder.",
        "**Paraphrased text:** the corpus is condensed and may omit nuances; it must not be treated as an authoritative copy of the law or standards.",
        "**Metric proxy:** answer correctness is token overlap with a gold span, not a human judgement; it under-credits correct paraphrases and over-credits long answers that happen to contain the span.",
        "**Optimistic tuning:** the TF-IDF blend and sentence-selection heuristics were adjusted after inspecting gold-set errors.",
        "**Future work:** hybrid dense + sparse retrieval with reciprocal-rank fusion; a cross-encoder re-ranker; an LLM-as-judge and RAGAS-style faithfulness scoring; OCR for scanned PDFs; multilingual (Hindi) support; scheduled corpus refresh with change alerts; and persistent storage with user authentication."])

    # ---- 13 Conclusion
    D.h("13. Conclusion", 1)
    D.p(f"The project delivers a complete, documented RAG prototype with a defensible evaluation: a safe offline baseline that finds the right document first for {pct(G['doc_hit@1'])} of 150 gold questions, "
        f"returns the correct fact in {pct(G['answer_correct_rate'])} of cases, refuses most out-of-scope questions, and always shows its evidence. The failure analysis shows where an LLM generator is expected to help, and the code is arranged so that the "
        "dense and LLM configurations can be measured with a single command each. The work follows the brief's ethical requirement that source chunks accompany every answer and that the system is not a substitute for legal or medical advice.")

    # ---- refs
    D.h("References", 1)
    refs = [
        "Lewis, P., Perez, E., Piktus, A., et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. Advances in Neural Information Processing Systems 33. arXiv:2005.11401.",
        "Salton, G. and Buckley, C. (1988). Term-weighting approaches in automatic text retrieval. Information Processing & Management, 24(5), 513-523.",
        "Karpukhin, V., Oguz, B., Min, S., et al. (2020). Dense Passage Retrieval for Open-Domain Question Answering. Proceedings of EMNLP 2020. arXiv:2004.04906.",
        "Reimers, N. and Gurevych, I. (2019). Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. Proceedings of EMNLP-IJCNLP 2019. arXiv:1908.10084.",
        "Ji, Z., Lee, N., Frieske, R., et al. (2023). Survey of Hallucination in Natural Language Generation. ACM Computing Surveys, 55(12).",
        "Liu, N. F., Lin, K., Hewitt, J., et al. (2024). Lost in the Middle: How Language Models Use Long Contexts. Transactions of the ACL, 12.",
        "National Institute of Standards and Technology (2023). Artificial Intelligence Risk Management Framework (AI RMF 1.0), NIST AI 100-1.",
        "Government of India (2023). The Digital Personal Data Protection Act, 2023. https://www.meity.gov.in",
        "LangChain documentation, https://python.langchain.com ; Chroma documentation, https://docs.trychroma.com ; Streamlit documentation, https://docs.streamlit.io",
        "Corpus sources: Constitution of India (https://legislative.gov.in); ministry portals such as https://pmkisan.gov.in, https://pmjay.gov.in, https://rti.gov.in; NIST (https://csrc.nist.gov); OSHA (https://www.osha.gov); FEMA (https://www.ready.gov); USDA (https://www.foodsafety.gov).",
        "HCLTech Career Shaper Programme Guidelines (provided by Programme Coordinator); mentor-validated references to be added: " + FILL]
    D.bullets(refs, style="List Number")

    # ---- appendices
    D.h("Appendix A. AI-tool disclosure", 1)
    D.p(AI_DISCLOSURE)
    D.h("Appendix B. Repository structure", 1)
    D.p("app.py (Streamlit UI) | ragqa/ (package) | scripts/ (build, evaluate, clean, documents) | tests/ | data/ (corpus, QA, dataset, manifest) | notebooks/EDA.ipynb | "
        "reports/ (documents, evaluation outputs, figures) | docs/ (diagrams, demo script) | requirements*.txt | .github/workflows/ci.yml | .streamlit/ | README.md.", size=10)
    D.h("Appendix C. Reproducing the results", 1)
    for cmd in ["pip install -r requirements-dev.txt", "python -m scripts.build_corpus && python -m scripts.make_qa", "python -m scripts.prepare_dataset",
                "python -m scripts.evaluate --sweep", "python -m scripts.build_notebook", "python -m pytest -q", "streamlit run app.py"]:
        D.p(cmd, size=10, after=1).runs[0].font.name = "Consolas"
    D.h("Appendix D. Deliverables checklist", 1)
    D.table(["Deliverable (from the brief)", "Where"],
            [["Completed project (working prototype)", "app.py + ragqa/ (run: streamlit run app.py)"],
             ["Project report, minimum 8 pages", "reports/Project_Report.docx (this document)"],
             ["EDA notebook with 3+ visualisations", "notebooks/EDA.ipynb (6 visualisations, executed)"],
             ["dataset.csv, cleaned and pre-processed", "data/dataset.csv (+ data/dataset_raw.csv)"],
             ["data_dictionary.csv, all columns described", "data_dictionary.csv"],
             ["Demo video (1-2 minutes) or live demo", "docs/demo_script.md (script and shot list; the video itself must be recorded by the learner)"],
             ["Signed Ethical Compliance Declaration", "reports/Ethical_Compliance_Declaration.docx (to be reviewed, ticked and signed by the learner)"],
             ["Project synopsis (template)", "reports/Project_Synopsis.docx (name, roll number, mentor and date to be filled in)"]],
            widths=[7.5, 8.7], caption="Mapping of deliverables to files")
    D.save(REP / "Project_Report.docx")


# ----------------------------------------------------------------------------- SYNOPSIS (fills the provided template)
def build_synopsis():
    d = Document(str(ROOT / "docs" / "Project_Synopsis_Template.docx"))
    P = d.paragraphs
    P[0].runs[0].text = "Project Synopsis"

    def fill(par, lines):
        run = par.runs[0]
        run.text = ""
        run.font.italic = False; run.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
        for i, line in enumerate(lines):
            if i:
                run.add_break()
            run.add_text(line)

    fill(P[5], ["Students and analysts who must find specific facts in long public documents such as the Constitution of India, welfare-scheme rules and technical standards "
                "currently scan PDFs by keyword, while generic chatbots answer fluently but may invent details. A wrong figure or article can mislead decisions, and in legal or health-related "
                "contexts it can cause harm. This project builds a RAG question-answering system that answers only from a 50-document public corpus (and user-uploaded files), "
                "shows the source chunk beside every answer, and refuses to answer when the evidence is weak."])
    fill(P[7], ["I built a Streamlit application and a Python package (ragqa) that splits documents into sentence-aware chunks, retrieves the most similar chunks and returns an answer with its "
                "source chunks, a confidence score and legal/medical notices. Retrieval uses ChromaDB through LangChain with local MiniLM or OpenAI embeddings, or an offline TF-IDF fallback; "
                "answers come from an OpenAI/Ollama LangChain chain or an extractive baseline that cannot hallucinate. An abstention threshold, a 1,000-pair Q&A set and an evaluation script "
                "measure and control quality."])
    fill(P[9], ["Figures below: (a) Data Flow Diagram Level 0, (b) system block diagram."])
    # add pictures to paragraph 9 itself (a then b)
    run = P[9].add_run(); run.add_break(); run.add_picture(str(DOCS / "architecture_dfd0.png"), width=Cm(12.5))
    run = P[9].add_run(); run.add_break(); run.add_picture(str(DOCS / "architecture_blocks.png"), width=Cm(15.0))
    fill(P[11], ["Frontend / UI: Streamlit (chat, source-chunk panel, feedback, corpus browser)",
                 "Backend / Logic: Python 3.11, ragqa package, pypdf, python-docx",
                 "Database: ChromaDB vector store (cosine) or in-memory TF-IDF index; CSV interaction log",
                 "AI/ML Framework: LangChain, ChromaDB, all-MiniLM-L6-v2 (ONNX) or OpenAI embeddings, OpenAI chat model or Ollama, scikit-learn",
                 "Visualisation: matplotlib, seaborn, scipy (EDA notebook); Streamlit metrics",
                 "Hardware: standard laptop/CPU; no GPU required"])
    fill(P[13], ["Dataset file name: dataset.csv (provided in project folder; extended and cleaned)",
                 f"Number of rows: {len(DF):,} (500 provided + 4,500 generated, same schema)",
                 f"Number of columns: {DF.shape[1]} (7 original + 7 derived)",
                 "Key columns: doc_id, doc_type, retrieval_score, query_type, answer_quality_score (also chunk_id, word_count)",
                 "Any additional data collected: 50-document public corpus (8,723 words), 1,000 Q&A pairs (150 hand-written + 850 synthetic), 30 out-of-scope questions, 1,000-row Q&A log in the brief's corpus schema"])
    fill(P[15], [f"Metric 1: Retrieval - correct document ranked first (Hit@1) / in top 3 (Hit@3) on 150 gold questions - {pct(G['doc_hit@1'])} / {pct(G['doc_hit@3'])} (MRR {G['doc_mrr@5']:.3f})",
                 f"Metric 2: Response time (retrieval + extractive answer, CPU) - mean {G['latency_ms_mean']:.1f} ms, p95 {G['latency_ms_p95']:.1f} ms (LLM latency not measured)",
                 f"Metric 3: Answer contains the gold span - {pct(G['answer_correct_rate'])}; out-of-scope questions refused - {pct(AB['out_of_scope_correctly_abstained'])} (in-scope wrongly refused {pct(AB['in_scope_wrongly_abstained'])})",
                 "Business value achieved: a cited, checkable answer in about a second instead of manual scanning, with built-in refusal and disclaimers. "
                 "Measured on the offline baseline over a small curated corpus; dense-embedding and LLM results are still to be run (see report Section 7.6)."])
    fill(P[17], ["Challenge 1: Time-sensitive facts (MGNREGA was replaced by VB-G RAM G on 1 July 2026) → Resolution: verified against current sources, rewrote the document, added as_of dates and a verify-with-official-source warning.",
                 "Challenge 2: Lexical retrieval confused \"money\" with a document on Money Bills → Resolution: blended word and character n-gram TF-IDF (hit@1 92.7% → 96.7%) and down-weighted title words in answer selection.",
                 "Challenge 3: Synthetic questions leaked answers (near-100% scores) → Resolution: reported them separately and used only the 150 hand-written questions for headline results.",
                 "Challenge 4: No network in the build environment → Resolution: lazy imports and an offline baseline so the pipeline is testable; dense/LLM paths clearly marked as not yet measured."])
    fill(P[19], ["1. Lewis, P. et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS 33. arXiv:2005.11401.",
                 "2. Karpukhin, V. et al. (2020). Dense Passage Retrieval for Open-Domain Question Answering. EMNLP 2020. arXiv:2004.04906.",
                 "3. Ji, Z. et al. (2023). Survey of Hallucination in Natural Language Generation. ACM Computing Surveys 55(12).",
                 "4. NIST (2023). AI Risk Management Framework 1.0 (NIST AI 100-1).",
                 "5. LangChain (https://python.langchain.com) and Chroma (https://docs.trychroma.com) documentation."])
    # template paragraph 21 (Ethical Compliance Declaration checkboxes + signature) is intentionally left for the learner
    # appendix
    d.add_page_break()
    h = d.add_paragraph(); r = h.add_run("Appendix - AI-tool disclosure"); r.bold = True; r.font.size = Pt(13)
    d.add_paragraph(AI_DISCLOSURE)
    d.save(str(REP / "Project_Synopsis.docx")); print("wrote", REP / "Project_Synopsis.docx")


# ----------------------------------------------------------------------------- DECLARATION
def build_declaration():
    D = Doc()
    D.p("HCLTech Career Shaper Programme  |  III Year Industry Project  |  2025-26", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    D.p("Ethical Compliance Declaration", bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    D.p("Project 72: RAG-Based Document Q&A System", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    D.table(["Field", "Details"], [["Learner name", FILL], ["Roll number", FILL], ["College / campus", FILL], ["Mentor name", FILL]], widths=[5, 11], size=10)
    D.p("**To the learner:** this form has been pre-filled with the evidence from the project but it is **not signed**. Read each statement, check that it is true for your submission, "
        "tick the box (replace the empty box with a check mark), then sign and date it. Do not sign a statement you have not verified.", size=10)
    D.h("I confirm that:", 2)
    D.table(["", "Statement", "Evidence in this project (for you to verify)"],
            [["☐", "All dataset usage complies with the source licence.", "Corpus = public government material (Government of India publications; US federal public-domain works), paraphrased, with source and licence note per document in data/corpus_manifest.csv. dataset.csv = the programme-supplied file plus synthetic rows generated for this project."],
             ["☐", "No personally identifiable information (PII) is stored without consent.", "No PII in corpus or datasets. App logs typed questions only to a local git-ignored CSV (logs/interactions.csv); uploads are processed in memory; users are advised not to enter personal data."],
             ["☐", "AI tools used for code/writing are disclosed in the Appendix.", "Appendix A below (also Appendix A of the report and the synopsis)."],
             ["☐", "This project complies with the DPDP Act 2023 where applicable.", "No personal data processed by design; if the OpenAI option is used, question and passage text is sent to that provider, which the user must accept; keys are kept in secrets, never in the repository."],
             ["☐", "The system displays the source chunk with every answer (brief, Section 8).", "app.py renders an expanded source-chunk panel for every answer, including refusals; unit test test_answer_always_ships_source_chunks."],
             ["☐", "The system is not presented as legal or medical advice.", "Permanent warning banner; automatic notices for personal legal/medical questions; limitations stated in README and report Section 9."]],
            widths=[1.0, 5.6, 9.6], size=9, caption="Declaration items")
    D.p("")
    D.p("Learner signature: _____________________________        Date: ____________________", size=11)
    D.p("")
    D.p("Mentor countersignature (if required): ____________________      Date: ______________", size=11)
    D.h("Appendix A. AI-tool disclosure", 2)
    D.p(AI_DISCLOSURE, size=10)
    D.h("Appendix B. Known limitations relevant to ethical use", 2)
    D.bullets(["RAG answers can still be wrong; the source chunks must be read.", "Corpus text is a condensed paraphrase and may be out of date; verify with the official source.",
               "Dense-retrieval and LLM configurations were not evaluated in the development environment.", "Monitoring-dataset scores are synthetic and do not describe real users or a real system."], style="List Bullet")
    D.save(REP / "Ethical_Compliance_Declaration.docx")


if __name__ == "__main__":
    REP.mkdir(exist_ok=True)
    build_report(); build_synopsis(); build_declaration()
